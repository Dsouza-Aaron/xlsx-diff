import openpyxl
import tkinter as tk
from tkinter import *
from tkinter import filedialog
from tkinter import ttk


# Custom class to highlight individual cells
class CellHighlightedTreeview(ttk.Treeview):
    def __init__(self, parent, **kwargs):
        super().__init__(parent, **kwargs)
        self.highlights = {}
        self.overlay_widgets = {}

        self.bind("<Configure>", lambda e: self.refresh_highlights())
        self.bind("<B1-Motion>", lambda e: self.refresh_highlights())
        self.bind("<MouseWheel>", lambda e: self.refresh_highlights())
        self.bind("<Button-4>", lambda e: self.refresh_highlights())
        self.bind("<Button-5>", lambda e: self.refresh_highlights())

    def highlight_cell(self, item_id, column, color="#ffcccc"):
        self.highlights[(item_id, column)] = color
        self.refresh_highlights()

    def clear_highlights(self):
        for widget in self.overlay_widgets.values():
            widget.destroy()
        self.overlay_widgets.clear()
        self.highlights.clear()

    def xview(self, *args):
        super().xview(*args)
        self.refresh_highlights()

    def yview(self, *args):
        super().yview(*args)
        self.refresh_highlights()

    def refresh_highlights(self):
        for widget in self.overlay_widgets.values():
            widget.place_forget()

        for (item_id, column), color in self.highlights.items():
            try:
                bbox = self.bbox(item_id, column)
                if bbox:
                    x, y, width, height = bbox

                    if (item_id, column) not in self.overlay_widgets:
                        lbl = tk.Label(self, anchor="w", padx=4, relief="flat")
                        lbl.bind(
                            "<Button-1>",
                            lambda e, i=item_id: self.selection_set(i),
                        )
                        self.overlay_widgets[(item_id, column)] = lbl

                    lbl = self.overlay_widgets[(item_id, column)]
                    val = self.set(item_id, column)

                    lbl.config(text=val, bg=color)
                    lbl.place(x=x, y=y, width=width, height=height)
                    lbl.lift()
            except Exception:
                pass


# GUI setup
root = Tk()
root.title("CompareExcel")
root.geometry("1920x1080")

# Top bar setup
top_bar = Frame(root, bg="#f0f2f5", height=45)
top_bar.pack(side=TOP, fill=X)
top_bar.pack_propagate(False)

file_1_adr = StringVar(value="File 1: None")
file_2_adr = StringVar(value="File 2: None")
sync_var = BooleanVar(value=True)

# Content frame using Grid layout
content_frame = Frame(root)
content_frame.pack(fill=BOTH, expand=True)

content_frame.rowconfigure(
    0, weight=1
)  # <--- Fixes vertical empty space/stretching
content_frame.columnconfigure(0, weight=1, uniform="group1")
content_frame.columnconfigure(1, weight=1, uniform="group1")
content_frame.columnconfigure(2, weight=0)

# Frame for first sheet
frame_1 = Frame(content_frame)
frame_1.grid(row=0, column=0, sticky="nsew")

# Frame for second sheet
frame_2 = Frame(content_frame)
frame_2.grid(row=0, column=1, sticky="nsew")

# Frame for log window
frame_3 = Frame(content_frame, width=250, bg="#e1e4e8")
frame_3.grid(row=0, column=2, sticky="nsew")
frame_3.pack_propagate(False)

# Spreadsheets setup
wb_1, sheet1 = None, None
wb_2, sheet2 = None, None
tree_items_1 = []
tree_items_2 = []
tree_1 = None
tree_2 = None

is_scrolling = False


# Scroll bar handlers
def scroll_y_1(*args):
    global is_scrolling
    if is_scrolling:
        return
    is_scrolling = True
    try:
        tree_1.yview(*args)
        if sync_var.get() and tree_2:
            tree_2.yview(*args)
    finally:
        is_scrolling = False


def scroll_y_2(*args):
    global is_scrolling
    if is_scrolling:
        return
    is_scrolling = True
    try:
        tree_2.yview(*args)
        if sync_var.get() and tree_1:
            tree_1.yview(*args)
    finally:
        is_scrolling = False


def scroll_x_1(*args):
    global is_scrolling
    if is_scrolling:
        return
    is_scrolling = True
    try:
        tree_1.xview(*args)
        if sync_var.get() and tree_2:
            tree_2.xview(*args)
    finally:
        is_scrolling = False


def scroll_x_2(*args):
    global is_scrolling
    if is_scrolling:
        return
    is_scrolling = True
    try:
        tree_2.xview(*args)
        if sync_var.get() and tree_1:
            tree_1.xview(*args)
    finally:
        is_scrolling = False


# Spreadsheets display, non-editable
def display_spreadsheet(parent_frame, file_path, file_num):
    global tree_items_1, tree_items_2, tree_1, tree_2

    for widget in parent_frame.winfo_children():
        widget.destroy()

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
    except Exception as e:
        Label(parent_frame, text=f"Error: {e}", fg="red").pack(
            padx=10, pady=10
        )
        return

    container = Frame(parent_frame)
    container.pack(fill=BOTH, expand=True)

    scroll_y = ttk.Scrollbar(container, orient=VERTICAL)
    scroll_x = ttk.Scrollbar(container, orient=HORIZONTAL)

    tree = CellHighlightedTreeview(
        container,
        yscrollcommand=scroll_y.set,
        xscrollcommand=scroll_x.set,
    )

    tree.tag_configure("mismatch", background="#ffcccc")

    if file_num == 1:
        scroll_y.config(command=scroll_y_1)
        scroll_x.config(command=scroll_x_1)
    else:
        scroll_y.config(command=scroll_y_2)
        scroll_x.config(command=scroll_x_2)

    scroll_y.pack(side=RIGHT, fill=Y)
    scroll_x.pack(side=BOTTOM, fill=X)
    tree.pack(side=LEFT, fill=BOTH, expand=True)

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return

    col_ids = [f"col_{i}" for i in range(len(rows[0]))]
    tree["columns"] = col_ids
    tree["show"] = "headings"

    for idx, col_id in enumerate(col_ids):
        header_text = str(rows[0][idx]) if rows[0][idx] is not None else ""
        tree.heading(col_id, text=header_text)
        tree.column(col_id, width=120, anchor="w")

    current_items = []
    for row in rows[1:]:
        values = [str(cell) if cell is not None else "" for cell in row]
        item_id = tree.insert("", END, values=values)
        current_items.append(item_id)

    if file_num == 1:
        tree_items_1 = current_items
        tree_1 = tree
    elif file_num == 2:
        tree_items_2 = current_items
        tree_2 = tree


# Get address of the sheets
def getfileaddress(file_adr, title_text, target_frame, file_num):
    global wb_1, sheet1, wb_2, sheet2
    filename = filedialog.askopenfilename(
        initialdir="C:/",
        title=title_text,
        filetypes=[("Excel Files", "*.xlsx")],
    )
    if filename:
        file_adr.set(filename)
        display_spreadsheet(target_frame, filename, file_num)

        if file_num == 1:
            wb_1 = openpyxl.load_workbook(filename, data_only=True)
            sheet1 = wb_1.active
        elif file_num == 2:
            wb_2 = openpyxl.load_workbook(filename, data_only=True)
            sheet2 = wb_2.active


# Primary function: prints log of all differences & highlights cells
def compare_xlsx():
    error_text.delete("1.0", END)
    if sheet1 is None or sheet2 is None:
        error_text.insert(END, "Please select both files first.\n")
        return

    if tree_1:
        tree_1.clear_highlights()
    if tree_2:
        tree_2.clear_highlights()

    for item in tree_items_1:
        tree_1.item(item, tags=())
    for item in tree_items_2:
        tree_2.item(item, tags=())

    max_r = max(sheet1.max_row, sheet2.max_row)
    max_c = max(sheet1.max_column, sheet2.max_column)

    differences_found = False

    for rowidx in range(1, max_r + 1):
        for colidx in range(1, max_c + 1):
            val1 = sheet1.cell(row=rowidx, column=colidx).value
            val2 = sheet2.cell(row=rowidx, column=colidx).value

            v1_str = "" if val1 is None else str(val1)
            v2_str = "" if val2 is None else str(val2)

            if v1_str != v2_str:
                differences_found = True
                cell_coord = (
                    openpyxl.utils.get_column_letter(colidx) + str(rowidx)
                )
                error_text.insert(
                    END,
                    f"Diff at {cell_coord}:\n  File 1: '{v1_str}'\n  File 2:"
                    f" '{v2_str}'\n\n",
                )

                # Skip header row (rowidx = 1) for tree item index
                if rowidx >= 2:
                    tree_row_index = rowidx - 2
                    col_id = f"col_{colidx - 1}"

                    if (
                        0 <= tree_row_index < len(tree_items_1)
                        and col_id in tree_1["columns"]
                    ):
                        tree_1.highlight_cell(
                            tree_items_1[tree_row_index], col_id, color="#ffcccc"
                        )
                    if (
                        0 <= tree_row_index < len(tree_items_2)
                        and col_id in tree_2["columns"]
                    ):
                        tree_2.highlight_cell(
                            tree_items_2[tree_row_index], col_id, color="#ffcccc"
                        )

    if not differences_found:
        error_text.insert(END, "No differences found!")


# GUI setup buttons
file_1_button = Button(
    top_bar,
    text="Select First File",
    command=lambda: getfileaddress(file_1_adr, "File 1", frame_1, 1),
    bg="#007acc",
    fg="white",
    relief=FLAT,
    padx=10,
)
file_1_button.pack(side=LEFT, padx=5, pady=8)

file_1_label = Label(
    top_bar, textvariable=file_1_adr, bg="#f0f2f5", fg="#333333", anchor="w"
)
file_1_label.pack(side=LEFT, padx=5)

file_2_button = Button(
    top_bar,
    text="Select Second File",
    command=lambda: getfileaddress(file_2_adr, "File_2", frame_2, 2),
    bg="#007acc",
    fg="white",
    relief=FLAT,
    padx=10,
)
file_2_button.pack(side=LEFT, padx=(15, 5), pady=8)

file_2_label = Label(
    top_bar, textvariable=file_2_adr, bg="#f0f2f5", fg="#333333", anchor="w"
)
file_2_label.pack(side=LEFT, padx=5)

execute_button = Button(
    top_bar,
    text="Execute",
    command=compare_xlsx,
    bg="#28a745",
    fg="white",
    activebackground="#218838",
    activeforeground="white",
    relief=FLAT,
    padx=12,
    pady=4,
    cursor="hand2",
)
execute_button.pack(side=LEFT, padx=10, pady=8)

toggle_container = Frame(top_bar, bg="#f0f2f5")
toggle_container.pack(side=LEFT, padx=10)

sync_label = Label(
    toggle_container,
    text="Scroll Sync",
    bg="#f0f2f5",
    fg="#333333",
    font=("Arial", 8),
)
sync_label.pack(anchor=W)

sync_check = Checkbutton(
    toggle_container,
    text="Enable Sync",
    variable=sync_var,
    bg="#f0f2f5",
    activebackground="#f0f2f5",
    relief=FLAT,
)
sync_check.pack(anchor=W)

error_container = Frame(frame_3, bg="#e1e4e8")
error_container.pack(fill=BOTH, expand=True)

text_container = Frame(error_container)
text_container.pack(fill=BOTH, expand=True)

text_scroll = ttk.Scrollbar(text_container, orient=VERTICAL)
text_scroll.pack(side=RIGHT, fill=Y)

error_text = Text(
    text_container,
    yscrollcommand=text_scroll.set,
    wrap=WORD,
    bg="white",
    fg="#333333",
    relief=FLAT,
)
error_text.pack(side=LEFT, fill=BOTH, expand=True)
text_scroll.config(command=error_text.yview)

root.mainloop()